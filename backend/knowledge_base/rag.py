"""Text extraction, chunking and retrieval for the portfolio-wide knowledge
base — the "RAG file" half of it. A document is turned into plain text,
split into overlapping chunks and saved as `<doc id>.rag.json` next to the
original. Each chunk also carries a sentence embedding (embed.py, local
all-MiniLM-L6-v2), and retrieval blends BM25 keyword ranking with cosine
similarity of those vectors — keywords catch exact names and numbers, vectors
catch paraphrases. If the embedding model can't load, retrieval falls back to
BM25 alone. The retrieved chunks are what get handed to Claude in
claude_client.py to actually answer with.
"""

import json
import math
import re
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree

import numpy as np

from . import embed

DOCUMENTS_ROOT = Path(__file__).resolve().parent.parent / 'documents'
KNOWLEDGE_BASE_DIR = DOCUMENTS_ROOT / 'knowledge_base'
UPLOADS_DIR = DOCUMENTS_ROOT / 'uploads'

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
# What can be turned into text for the knowledge base. Upload-only documents
# aren't limited to these — they're just stored for people to download.
INDEXABLE_EXTENSIONS = {'.pdf', '.docx', '.xlsx', '.xlsm', '.txt', '.md', '.csv', '.json'}
BLOCKED_EXTENSIONS = {'.exe', '.bat', '.cmd', '.com', '.msi', '.scr', '.js', '.vbs', '.ps1', '.sh', '.dll'}

CHUNK_TARGET_CHARS = 1000
CHUNK_OVERLAP_CHARS = 180

_STOPWORDS = {
    'the', 'and', 'for', 'are', 'was', 'with', 'that', 'this', 'from', 'have', 'has', 'not', 'but', 'you',
    'que', 'los', 'las', 'del', 'con', 'una', 'por', 'para', 'como', 'los', 'sus', 'más', 'pero',
}


def ensure_dir(path: Path) -> Path:
    path.mkdir(parents=True, exist_ok=True)
    return path


def rag_path(doc_id: str) -> Path:
    return ensure_dir(KNOWLEDGE_BASE_DIR) / f'{doc_id}.rag.json'


# ---------------------------------------------------------------- extraction


def _extract_pdf(path: Path) -> str:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        pages = []
        for index in range(len(pdf)):
            page = pdf[index]
            text_page = page.get_textpage()
            pages.append(text_page.get_text_range())
            text_page.close()
            page.close()
        return '\n\n'.join(pages)
    finally:
        pdf.close()


def _extract_docx(path: Path) -> str:
    namespace = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read('word/document.xml'))
    paragraphs = []
    for paragraph in root.iter(f'{namespace}p'):
        text = ''.join(node.text or '' for node in paragraph.iter(f'{namespace}t'))
        if text.strip():
            paragraphs.append(text)
    return '\n\n'.join(paragraphs)


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        lines = [f'Sheet: {sheet.title}']
        for row in sheet.iter_rows(values_only=True):
            cells = ['' if value is None else str(value) for value in row]
            if any(cell.strip() for cell in cells):
                lines.append(' | '.join(cells).strip())
        parts.append('\n'.join(lines))
    workbook.close()
    return '\n\n'.join(parts)


def extract_text(path: Path) -> str:
    """Plain text of a supported file. Raises ValueError with a message
    fit to show a user if the type isn't supported or nothing readable
    comes out (e.g. a scanned, image-only PDF)."""
    extension = path.suffix.lower()
    if extension not in INDEXABLE_EXTENSIONS:
        raise ValueError(
            f'"{extension or "this file type"}" can\'t be added to the knowledge base — supported types are '
            + ', '.join(sorted(INDEXABLE_EXTENSIONS))
            + '.'
        )
    try:
        if extension == '.pdf':
            text = _extract_pdf(path)
        elif extension == '.docx':
            text = _extract_docx(path)
        elif extension in ('.xlsx', '.xlsm'):
            text = _extract_xlsx(path)
        else:
            text = path.read_bytes().decode('utf-8', errors='replace')
    except Exception as error:  # a corrupt/encrypted file is a user problem, not a server one
        raise ValueError(f'Could not read this file: {error}') from error

    text = re.sub(r'[ \t]+\n', '\n', text).strip()
    if not text:
        raise ValueError('No readable text was found in this file (a scanned, image-only PDF for example).')
    return text


# ------------------------------------------------------------------ chunking


def chunk_text(text: str) -> list[str]:
    """Paragraph-aware chunks of ~CHUNK_TARGET_CHARS, each starting with the
    tail of the previous one so a sentence split across a boundary is still
    findable from either side."""
    paragraphs = [part.strip() for part in re.split(r'\n\s*\n', text) if part.strip()]
    pieces: list[str] = []
    for paragraph in paragraphs:
        if len(paragraph) <= CHUNK_TARGET_CHARS:
            pieces.append(paragraph)
            continue
        # A very long paragraph (a whole page of a PDF with no blank lines)
        # is cut on sentence-ish boundaries instead.
        sentences = re.split(r'(?<=[.!?])\s+', paragraph)
        current = ''
        for sentence in sentences:
            while len(sentence) > CHUNK_TARGET_CHARS:
                pieces.append(sentence[:CHUNK_TARGET_CHARS])
                sentence = sentence[CHUNK_TARGET_CHARS:]
            if current and len(current) + len(sentence) + 1 > CHUNK_TARGET_CHARS:
                pieces.append(current)
                current = sentence
            else:
                current = f'{current} {sentence}'.strip()
        if current:
            pieces.append(current)

    chunks: list[str] = []
    current = ''
    for piece in pieces:
        if current and len(current) + len(piece) + 2 > CHUNK_TARGET_CHARS:
            chunks.append(current)
            overlap = current[-CHUNK_OVERLAP_CHARS:]
            current = f'{overlap}\n\n{piece}'
        else:
            current = f'{current}\n\n{piece}'.strip()
    if current:
        chunks.append(current)
    return chunks


def build_rag_file(doc_id: str, name: str, source: str, text: str) -> int:
    chunks = chunk_text(text)
    payload = {
        'doc_id': doc_id,
        'name': name,
        'source': source,
        'chunks': [{'index': index, 'text': chunk} for index, chunk in enumerate(chunks)],
    }
    _attach_embeddings(payload)
    rag_path(doc_id).write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
    return len(chunks)


def _attach_embeddings(data: dict) -> bool:
    """Embed any chunk that has no vector yet. Returns True if anything was
    added. Best effort: a missing/broken model leaves the file BM25-only."""
    missing = [chunk for chunk in data.get('chunks', []) if not chunk.get('embedding')]
    if not missing:
        return False
    try:
        vectors = embed.embed_texts([chunk['text'] for chunk in missing])
    except Exception:
        return False
    for chunk, vector in zip(missing, vectors):
        chunk['embedding'] = embed.encode_vector(vector)
    data['embedding_model'] = embed.MODEL_ID
    return True


def delete_rag_file(doc_id: str) -> None:
    rag_path(doc_id).unlink(missing_ok=True)


# ----------------------------------------------------------------- retrieval


def _tokens(text: str) -> list[str]:
    return [token for token in re.findall(r'\w+', text.lower()) if len(token) > 1 and token not in _STOPWORDS]


BM25_WEIGHT = 0.4
VECTOR_WEIGHT = 0.6
MIN_COSINE = 0.2  # below this a chunk with no keyword hit is just noise


def search(doc_ids: list[str], query: str, top_k: int = 6) -> list[dict]:
    """Hybrid retrieval over every chunk of the given documents: BM25 blended
    with cosine similarity of chunk/query embeddings. Returns the best chunks,
    best first, each tagged with the document it came from. Files indexed
    before embeddings existed are embedded (and saved) on first search."""
    query_terms = _tokens(query)

    entries = []
    for doc_id in doc_ids:
        path = rag_path(doc_id)
        if not path.exists():
            continue
        data = json.loads(path.read_text(encoding='utf-8'))
        if _attach_embeddings(data):
            path.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
        for chunk in data.get('chunks', []):
            vector = embed.decode_vector(chunk['embedding']) if chunk.get('embedding') else None
            entries.append((data['doc_id'], data['name'], chunk['index'], chunk['text'], _tokens(chunk['text']), vector))
    if not entries:
        return []

    query_vector = None
    if any(entry[5] is not None for entry in entries) and query.strip():
        try:
            query_vector = embed.embed_text(query)
        except Exception:
            query_vector = None
    if not query_terms and query_vector is None:
        return []

    total = len(entries)
    average_length = sum(len(entry[4]) for entry in entries) / total or 1
    document_frequency: Counter = Counter()
    for entry in entries:
        for term in set(entry[4]):
            document_frequency[term] += 1

    k1, b = 1.5, 0.75
    rows = []
    for doc_id, name, index, text, tokens, vector in entries:
        counts = Counter(tokens)
        keyword = 0.0
        for term in set(query_terms):
            frequency = counts.get(term, 0)
            if not frequency:
                continue
            idf = math.log(1 + (total - document_frequency[term] + 0.5) / (document_frequency[term] + 0.5))
            keyword += idf * (frequency * (k1 + 1)) / (frequency + k1 * (1 - b + b * len(tokens) / average_length))
        cosine = float(np.dot(query_vector, vector)) if query_vector is not None and vector is not None else 0.0
        rows.append((doc_id, name, index, text, keyword, cosine))

    top_keyword = max((row[4] for row in rows), default=0.0)
    scored = []
    for doc_id, name, index, text, keyword, cosine in rows:
        if keyword <= 0 and cosine < MIN_COSINE:
            continue
        if query_vector is None:
            score = keyword
        else:
            score = BM25_WEIGHT * (keyword / top_keyword if top_keyword else 0.0) + VECTOR_WEIGHT * max(cosine, 0.0)
        scored.append({'doc_id': doc_id, 'doc_name': name, 'chunk_index': index, 'text': text, 'score': round(score, 4)})

    scored.sort(key=lambda item: item['score'], reverse=True)
    return scored[:top_k]
