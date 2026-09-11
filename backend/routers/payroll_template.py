"""
Generates the fillable payroll workbook for the "Download format" workflow:
download it, fill it in, and pick it back up via "Upload filled format" — a
plain download/upload round trip. This app and the spreadsheet editor don't
have to be on the same machine, and there's no background file watch to keep
alive across page navigation.

One row per position — Level drives compensation, so there's no per-employee
roster or per-year salary entry here. The Level reference table sits beside
the data as a lookup guide, in the same "Payroll" sheet rather than a
separate tab.
"""

import io

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from sqlalchemy.orm import Session

from ..database import get_db
from .. import models
from .payroll import _ensure_payroll_record, _get_effective_year_salary, _get_year_payroll_level, _parent_map

router = APIRouter()

# Column A is an unlabeled row number (a plain reading aid — never read back
# on import). "Compensation (standard)" is always the value implied by
# "Level" via the reference table on the right, filled in for reference only.
# "Compensation (custom)" is the one column that actually overrides a
# position's salary on import — pre-filled here only when the position's
# current salary doesn't match its level's standard figure, so re-downloading
# an already-customized position shows the user their existing override
# instead of silently looking blank (and re-uploading unedited is a no-op).
HEADER_STATIC = ['location', 'Position', 'Level', 'Compensation (standard)', 'Compensation (custom)', 'Area', 'Subordinated To']
REFERENCE_HEADER = ['Level', 'Yearly', 'Percentage', 'Monthly']
REFERENCE_START_COLUMN = 10  # column J, leaving column I blank as a spacer
REFERENCE_FILL = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')


def _build_workbook(db: Session, company_id: str) -> Workbook:
    nodes = db.query(models.OrgChartNode).filter_by(company_id=company_id).all()
    edges = db.query(models.OrgChartEdge).filter_by(company_id=company_id).all()
    parents = _parent_map(edges)
    node_by_id = {node.id: node for node in nodes}
    payroll_levels = db.query(models.PayrollLevel).order_by(models.PayrollLevel.sort_order).all()
    yearly_by_level = {level.level: level.yearly for level in payroll_levels}

    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Payroll'

    # Column A (the row number) is intentionally left header-less, matching
    # the row-number column already used for readability elsewhere.
    for column, header in enumerate(HEADER_STATIC, start=2):
        sheet.cell(row=1, column=column, value=header).font = sheet.cell(row=1, column=column).font.copy(bold=True)

    sorted_nodes = sorted(nodes, key=lambda item: (item.sort_index or 0.0, item.office_name))
    for row_offset, node in enumerate(sorted_nodes):
        row_index = row_offset + 2
        parent_id = parents.get(node.id)
        parent_name = node_by_id[parent_id].office_name if parent_id in node_by_id else ''
        record = _ensure_payroll_record(db, node)
        level = _get_year_payroll_level(db, record, 0)
        standard_salary = yearly_by_level.get(level) if level else None
        current_salary = _get_effective_year_salary(db, record, 0)
        custom_salary = current_salary if standard_salary is not None and abs(current_salary - standard_salary) > 0.01 else None

        sheet.cell(row=row_index, column=1, value=row_offset + 1)
        sheet.cell(row=row_index, column=2, value=node.location or '')
        sheet.cell(row=row_index, column=3, value=node.office_name)
        sheet.cell(row=row_index, column=4, value=level or '')
        sheet.cell(row=row_index, column=5, value=standard_salary)
        sheet.cell(row=row_index, column=6, value=custom_salary)
        sheet.cell(row=row_index, column=7, value=node.area or '')
        sheet.cell(row=row_index, column=8, value=parent_name)

    for column_offset, header in enumerate(REFERENCE_HEADER):
        column = REFERENCE_START_COLUMN + column_offset
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = cell.font.copy(bold=True)
        cell.fill = REFERENCE_FILL

    for row_index, level in enumerate(payroll_levels, start=2):
        values = [level.level, level.yearly, level.percentage, level.monthly]
        for column_offset, value in enumerate(values):
            cell = sheet.cell(row=row_index, column=REFERENCE_START_COLUMN + column_offset, value=value)
            cell.fill = REFERENCE_FILL

    instructions = wb.create_sheet('Instructions')
    for line in [
        'How this file works',
        '',
        '- One row per position. Column A is just a row number for readability.',
        '- "Level" must exactly match a level from the reference table on the right (columns J-M) —',
        '  that\'s what sets "Compensation (standard)" for every projection year. Leave "Level" blank',
        '  to leave the position\'s compensation unchanged.',
        '- "Compensation (standard)" is filled in automatically from "Level" — it\'s for reference only',
        '  and is not read back in; edit "Level" instead to change it.',
        '- "Compensation (custom)" overrides the standard, level-based figure for this one position.',
        '  Leave it blank to use the standard figure; fill it in to pay this position something',
        '  different from its level.',
        '- "Subordinated To" must exactly match another row\'s "Position" text (or "Board of Directors", or',
        '  blank for the top of the chart).',
        '- Save the file, then use "Upload filled format" in the app to bring your changes in.',
    ]:
        instructions.append([line])

    sheet.column_dimensions['A'].width = 6
    sheet.column_dimensions['B'].width = 12
    sheet.column_dimensions['C'].width = 28
    sheet.column_dimensions['D'].width = 10
    sheet.column_dimensions['E'].width = 16
    sheet.column_dimensions['F'].width = 16
    sheet.column_dimensions['G'].width = 22
    sheet.column_dimensions['H'].width = 22
    sheet.column_dimensions['I'].width = 4
    sheet.column_dimensions['J'].width = 10
    sheet.column_dimensions['K'].width = 12
    sheet.column_dimensions['L'].width = 12
    sheet.column_dimensions['M'].width = 12

    return wb


@router.get('/companies/{company_id}/payroll-template')
def download_payroll_template(company_id: str, db: Session = Depends(get_db)):
    company = db.query(models.Company).filter_by(id=company_id).first()
    if not company:
        raise HTTPException(status_code=404, detail='Company not found')

    workbook = _build_workbook(db, company_id)
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{company_id}-payroll-template.xlsx"'},
    )
